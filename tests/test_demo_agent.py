"""Real MAF orchestration against a deterministic reverse-link fixture; no remote inference in tests."""
import asyncio
import io
import json
import re
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "demo" / "frontend"))
from agent import agent_response
from app import Config, create_app
from hub import Gate
from workspace import Workspaces, Workspace, calculate, extract, number, safe_preview, table_rows
sys.path.pop(0)


class FilesTests(unittest.TestCase):
    def test_chat_model_selector_uses_reserved_content_row(self):
        static = ROOT / "demo" / "frontend" / "static"
        html = (static / "index.html").read_text()
        css = (static / "chat.css").read_text()
        js = (static / "chat.js").read_text()
        header = html.split("<header", 1)[1].split("</header>", 1)[0]
        scope = html.split('<div id="model-scope"', 1)[1].split("</div>", 1)[0]
        self.assertNotIn('id="model-select"', header)
        self.assertEqual(html.count('id="model-select"'), 1)
        self.assertIn('for="model-select"', scope)
        self.assertIn('aria-describedby="model-note"', scope)
        self.assertIn("모델 변경 시 새 대화가 시작됩니다.", scope)
        self.assertIn('value="fp8"', scope)
        self.assertIn('value="nvfp4"', scope)
        selector = css.split(".model-select {", 1)[1].split("}", 1)[0]
        for rule in ("width:12rem", "min-width:12rem", "height:2.145rem",
                     "padding:.35rem .85rem", "font-size:.85rem", "border-radius:8px"):
            self.assertIn(rule, selector)
        self.assertIn(".model-scope.inactive { visibility:hidden; }", css)
        self.assertIn("max-width:var(--chat-width)", css.split(".model-scope {", 1)[1].split("}", 1)[0])
        self.assertIn('ui.modelScope.classList.toggle("inactive", mode !== "chat");', js)
        self.assertIn('"model-scope"', js)

    def test_even_desktop_samples_and_web_iq_tool_enabling(self):
        js = (ROOT / "demo" / "frontend" / "static" / "chat.js").read_text()
        css = (ROOT / "demo" / "frontend" / "static" / "chat.css").read_text()
        examples = js.split("const EXAMPLES = [", 1)[1].split("];", 1)[0]
        samples = re.findall(r'\{ label: "([^"]+)", text: "([^"]+)", tools: true \}', examples)
        self.assertEqual(len(samples), 8)
        self.assertEqual(len(samples) % 2, 0)
        self.assertEqual(sum(label.startswith("Web IQ") for label, _ in samples), 4)
        for word in ("웹 검색", "뉴스 도구", "금융 도구", "장소 도구", "005930", "미확인", "문서·CSV"):
            self.assertIn(word, examples)
        self.assertIn("if (example.tools) ui.tools.checked = true;", js)
        self.assertIn("grid-template-columns: repeat(2, minmax(0, 1fr))", css)
        self.assertIn("@media (min-width: 1200px) { .suggestions { grid-template-columns: repeat(4, minmax(0, 1fr)); } }", css)
        self.assertNotIn("repeat(auto-fit", css.split(".suggestions {", 1)[1].split("}", 1)[0])

    def test_desktop_shared_width_native_typography_and_visible_download(self):
        css = (ROOT / "demo" / "frontend" / "static" / "chat.css").read_text()
        js = (ROOT / "demo" / "frontend" / "static" / "chat.js").read_text()
        self.assertIn("--chat-width: 1200px", css)
        self.assertIn(".messages > *, .dock > *", css)
        self.assertIn(".messages, .dock { scrollbar-gutter: stable both-edges; }", css)
        self.assertIn("font-size: 18px", css)
        self.assertIn("width: 42px; height: 42px", css)
        self.assertNotIn("zoom:", css)
        for selector in (":root", ".app", ".messages", ".dock", ".composer"):
            style = css.split(selector + " {", 1)[1].split("}", 1)[0]
            self.assertNotIn("transform: scale", style)
        self.assertIn("summary.append", js)
        self.assertIn("action.state === \"success\"", js)

    def test_exact_and_bounded_calculator(self):
        self.assertEqual(calculate("2^64-1"), "18446744073709551615")
        self.assertEqual(calculate("0.1+0.2"), "3/10")
        self.assertEqual(calculate("1/3"), "1/3")
        for expression in ("__import__('os')", "2^100000", "1e999999999", "True+1", "1/0"):
            with self.subTest(expression=expression), self.assertRaises((ValueError, ZeroDivisionError)):
                calculate(expression)

    def test_isolation_paths_and_zip_atomicity(self):
        a, b = Workspace(), Workspace()
        a.put("src/main.py", b"print('hello')", original=True)
        with self.assertRaises(ValueError):
            b.get("src/main.py")
        for name in ("../x.py", "/x.py", "a\\x.py", "C:x.py", "a/../x.py", "x.exe", "a//x.py"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                a.put(name, b"no")
        zipped = io.BytesIO()
        with zipfile.ZipFile(zipped, "w") as archive:
            archive.writestr("valid.py", "ok")
            archive.writestr("../escape.py", "bad")
        with self.assertRaises(ValueError):
            a.upload("project.zip", zipped.getvalue())
        self.assertNotIn("valid.py", a.files)
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.txt", "a" * 1_000_000)
        with self.assertRaises(ValueError):
            a.upload("bomb.zip", bomb.getvalue())

    def test_extraction_and_tables(self):
        from docx import Document
        from openpyxl import Workbook
        doc = Document()
        doc.add_paragraph("A.X sample")
        data = io.BytesIO()
        doc.save(data)
        self.assertIn("[paragraph 1] A.X sample", extract("sample.docx", data.getvalue()))
        self.assertIn("[line 1] hello", extract("sample.txt", b"hello"))
        self.assertEqual(table_rows("a.csv", b"region,sales\nSeoul,20\n"), [["region", "sales"], ["Seoul", "20"]])
        book = Workbook()
        book.active.append(["value"])
        book.active.append(["=1+1"])
        data = io.BytesIO()
        book.save(data)
        with self.assertRaises(ValueError):
            table_rows("a.xlsx", data.getvalue())
        self.assertEqual(str(number("0.1")), "1/10")
        for value in ("1e999999999", "=1+1", "9" * 257):
            with self.assertRaises(ValueError):
                number(value)

    def test_preview_no_active_or_url_content(self):
        text = safe_preview(b'<meta http-equiv="refresh" content="0;url=https://bad">'
                            b'<script>fetch("/api/admin")</script><iframe src="https://bad"></iframe>'
                            b'<a href="https://bad">link</a><img src="https://bad"><p onclick="alert(1)">hello</p>')
        self.assertNotIn("<script", text)
        self.assertNotIn("<meta", text)
        self.assertNotIn("<iframe", text)
        self.assertNotIn("href=", text)
        self.assertNotIn("src=", text)
        self.assertNotIn("onclick=", text)

    def test_expiry_capacity(self):
        manager = Workspaces()
        space = manager.create()
        space.touched -= manager.ttl + 1
        with self.assertRaises(ValueError):
            manager.get(space.token)
        self.assertEqual(len(manager.items), 0)
        for _ in range(16):
            manager.create()
        with self.assertRaises(ValueError):
            manager.create()


class FixtureLink:
    def __init__(self, name="calculator", arguments=None, hold=False):
        self.streams, self.requests, self.cancelled = {}, [], []
        self.name, self.arguments, self.hold = name, arguments or {"expression": "2^64-1"}, hold

    def send(self, message):
        if message["type"] == "cancel":
            self.cancelled.append(message["id"])
            return
        self.requests.append(message["body"])
        queue = self.streams[message["id"]]
        queue.put_nowait({"type": "head", "status": 200})
        if self.hold:
            return
        if len(self.requests) == 1:
            choices = [
                {"delta": {"reasoning_content": "real fixture reasoning"}},
                {"delta": {"tool_calls": [{"index": 0, "id": "call_1", "function": {
                    "name": self.name, "arguments": json.dumps(self.arguments)[:5]}}]}},
                {"delta": {"tool_calls": [{"index": 0, "function": {
                    "arguments": json.dumps(self.arguments)[5:]}}]}, "finish_reason": "tool_calls"},
            ]
        else:
            choices = [{"delta": {"content": "Final A.X answer"}, "finish_reason": "stop"}]
        for choice in choices:
            queue.put_nowait({"type": "data", "chunk": "data: " + json.dumps({"choices": [choice]}) + "\n\n"})
        queue.put_nowait({"type": "end"})


class FixtureHub:
    def __init__(self, link):
        self.link = link

    def ready_link(self):
        return self.link


class AgentTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_maf_invokes_tool_and_replays_reasoning(self):
        link, space, gate = FixtureLink(), Workspace(), Gate(1, 2)
        space.busy = True
        response = agent_response(FixtureHub(link), gate, {
            "model": "axk2", "stream": True, "messages": [{"role": "user", "content": "calculate"}],
            "chat_template_kwargs": {"enable_thinking": True}, "max_tokens": 512}, space, True)
        output = b"".join([chunk async for chunk in response.body_iterator]).decode()
        self.assertIn("18446744073709551615", output)
        self.assertIn('"state": "pending"', output)
        self.assertIn('"state": "running"', output)
        self.assertIn('"state": "success"', output)
        self.assertIn("Final A.X answer", output)
        self.assertEqual(len(link.requests), 2)
        replay = link.requests[1]["messages"]
        self.assertEqual(next(m for m in replay if m.get("tool_calls"))["reasoning_content"], "real fixture reasoning")
        self.assertIn("18446744073709551615", next(m["content"] for m in replay if m["role"] == "tool"))
        self.assertFalse(space.busy)
        self.assertEqual(gate.active, 0)
        self.assertNotIn("get_weather", str(link.requests))

    async def test_unconfigured_execution_is_real_tool_error(self):
        link, space = FixtureLink("run_tests", {"command": "pytest"}), Workspace()
        response = agent_response(FixtureHub(link), Gate(1, 2), {
            "messages": [{"role": "user", "content": "run tests"}], "model": "axk2"}, space, True)
        output = b"".join([chunk async for chunk in response.body_iterator]).decode()
        self.assertIn('"state": "error"', output)
        self.assertIn("No tests were run", output)
        self.assertNotIn('"state": "success"', output)

    async def test_generated_python_saved_action_and_exact_authorized_download(self):
        with tempfile.TemporaryDirectory() as folder:
            app = create_app(Config(data=Path(folder), aml_dir=ROOT / "aml", session_secret="test", open_demo=True),
                             start_supervisor=False)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                token = (await client.post("/api/workspace")).json()["token"]
                other = (await client.post("/api/workspace")).json()["token"]
                content = "def binary_search(items, target):\n    return -1\n"
                link = FixtureLink("write_file", {"path": "binary_search.py", "content": content})
                app.state.hub.ready_link = lambda: link
                capabilities = (await client.get("/api/capabilities")).json()
                self.assertEqual(capabilities["web_iq"]["state"], "unavailable")
                response = await client.post("/api/agent", headers={"X-AX-Workspace": token}, json={
                    "messages": [{"role": "user", "content": "write Python"}], "tools": True})
                self.assertEqual(response.status_code, 200)
                output = response.text
                actions = [json.loads(line[6:]) for line in output.splitlines()
                           if line.startswith("data: ") and '"state": "success"' in line]
                result = actions[0]["result"]
                self.assertEqual(result["artifact"], {"kind": "download", "path": "binary_search.py"})
                self.assertEqual(actions[0]["arguments"]["content"], content)
                downloaded = await client.get("/api/workspace/download?path=binary_search.py",
                                               headers={"X-AX-Workspace": token})
                self.assertEqual(downloaded.content, content.encode())
                self.assertIn('filename="binary_search.py"', downloaded.headers["content-disposition"])
                self.assertEqual(downloaded.headers["content-type"], "application/octet-stream")
                denied = await client.get("/api/workspace/download?path=binary_search.py",
                                           headers={"X-AX-Workspace": other})
                self.assertEqual(denied.status_code, 404)

    async def test_cancellation_releases_gate_and_link(self):
        link, space, gate = FixtureLink(hold=True), Workspace(), Gate(1, 2)
        response = agent_response(FixtureHub(link), gate, {
            "messages": [{"role": "user", "content": "hold"}], "model": "axk2"}, space, True)
        task = asyncio.create_task(self.consume(response))
        for _ in range(100):
            if link.requests:
                break
            await asyncio.sleep(.01)
        task.cancel()
        with self.assertRaises(asyncio.CancelledError):
            await task
        self.assertEqual(gate.active, 0)
        self.assertEqual(len(link.cancelled), 1)
        self.assertFalse(link.streams)
        self.assertFalse(space.busy)

    async def consume(self, response):
        async for _ in response.body_iterator:
            pass

    async def test_http_workspace_and_artifact_capability(self):
        with tempfile.TemporaryDirectory() as folder:
            app = create_app(Config(data=Path(folder), aml_dir=ROOT / "aml", session_secret="test", open_demo=True),
                             start_supervisor=False)
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                a = (await client.post("/api/workspace")).json()["token"]
                b = (await client.post("/api/workspace")).json()["token"]
                response = await client.post("/api/workspace/upload?name=index.html",
                                             headers={"X-AX-Workspace": a}, content=b"<h1>OK</h1><script>bad()</script>")
                self.assertEqual(response.status_code, 200)
                denied = await client.get("/api/workspace/download?path=index.html", headers={"X-AX-Workspace": b})
                self.assertEqual(denied.status_code, 404)
                denied = await client.get("/api/workspace/download?path=index.html")
                self.assertEqual(denied.status_code, 410)
                preview = await client.get("/api/workspace/preview?path=index.html", headers={"X-AX-Workspace": a})
                self.assertNotIn("<script", preview.text)
                download = await client.get("/api/workspace/download?path=index.html", headers={"X-AX-Workspace": a})
                self.assertIn("attachment", download.headers["content-disposition"])
                self.assertEqual(download.headers["content-type"], "application/octet-stream")


if __name__ == "__main__":
    unittest.main()
