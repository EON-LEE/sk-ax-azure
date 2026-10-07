"""Bounded, capability-addressed chat files. No uploaded or generated code is executed."""
import ast
import csv
import difflib
import io
import json
import re
import secrets
import stat
import time
import zipfile
from datetime import datetime
from fractions import Fraction
from pathlib import PurePosixPath
from zoneinfo import ZoneInfo

LIMIT = 8 * 1024 * 1024
TEXT_LIMIT = 200_000
EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx", ".html", ".css", ".js",
              ".py", ".json", ".yml", ".yaml", ".toml", ".xml", ".sql", ".ts"}


def safe_preview(data):
    import bleach
    # No URL-bearing attributes, active elements, CSS, SVG, meta refresh, links or forms.
    return bleach.clean(data.decode("utf-8"), tags={
        "h1", "h2", "h3", "h4", "p", "div", "span", "section", "article", "header", "footer", "main",
        "ul", "ol", "li", "pre", "code", "strong", "em", "b", "i", "br", "hr",
        "table", "thead", "tbody", "tr", "td", "th", "blockquote",
    }, attributes={}, protocols=[], strip=True)


def filename(name):
    if (not isinstance(name, str) or len(name) > 160 or "\\" in name or "\x00" in name
            or any(ord(c) < 32 for c in name) or ":" in name):
        raise ValueError("Invalid file path")
    path = PurePosixPath(name)
    if path.is_absolute() or any(p in ("", ".", "..") for p in name.split("/")):
        raise ValueError("Only relative project paths are allowed")
    if path.suffix.lower() not in EXTENSIONS:
        raise ValueError("Unsupported file type")
    return str(path)


def unzip(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 200 or sum(i.file_size for i in infos) > LIMIT:
            raise ValueError("Archive exceeds expanded file/count limits")
        for info in infos:
            if (info.flag_bits & 1 or stat.S_ISLNK(info.external_attr >> 16)
                    or info.file_size > LIMIT or info.file_size > max(1, info.compress_size) * 100):
                raise ValueError("Encrypted, linked or highly compressed archive is not allowed")
            path = PurePosixPath(info.filename)
            if (path.is_absolute() or ".." in path.parts or "\\" in info.filename
                    or ":" in info.filename):
                raise ValueError("Unsafe archive member path")
        return {i.filename: archive.read(i) for i in infos if not i.is_dir()}


def extract(name, data):
    suffix = PurePosixPath(name).suffix.lower()
    if suffix in {".docx", ".xlsx"}:
        unzip(data)  # expansion limits before Office parsers allocate
    if suffix == ".pdf":
        from pdf_documents import all_text
        return all_text(data, name)
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [f"[paragraph {i}] {p.text}" for i, p in enumerate(doc.paragraphs, 1)]
        for i, table in enumerate(doc.tables, 1):
            parts.extend(f"[table {i}, row {j}] " + " | ".join(c.text for c in row.cells)
                         for j, row in enumerate(table.rows, 1))
        text = "\n".join(parts)
    elif suffix == ".xlsx":
        return json.dumps(table_rows(name, data), ensure_ascii=False)
    else:
        text = data.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError("Binary files are not supported")
        text = "\n".join(f"[line {i}] {line}" for i, line in enumerate(text.splitlines(), 1))
    if len(text) > TEXT_LIMIT:
        raise ValueError("Extracted text exceeds limit")
    return text


def table_rows(name, data):
    if name.lower().endswith(".csv"):
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        result = []
        for row in rows:
            if len(row) > 50 or len(result) >= 2000 or sum(map(len, row)) > 20_000:
                raise ValueError("Table limit: 2000 rows, 50 columns, 20000 characters/row")
            result.append(row)
        return result
    if not name.lower().endswith(".xlsx"):
        raise ValueError("Use a CSV or XLSX file")
    from openpyxl import load_workbook
    unzip(data)
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book.worksheets[0]
        if sheet.max_row > 2000 or sheet.max_column > 50:
            raise ValueError("Table limit: 2000 rows and 50 columns")
        result = []
        for row in sheet.iter_rows():
            if any(cell.data_type == "f" for cell in row):
                raise ValueError("Spreadsheet formulas are not evaluated; upload values-only data")
            values = ["" if c.value is None else str(c.value) for c in row]
            if sum(map(len, values)) > 20_000 or len(result) >= 2000 or len(values) > 50:
                raise ValueError("Table limits exceeded")
            result.append(values)
        return result
    finally:
        book.close()


def number(text):
    if len(text) > 256 or not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:/[+-]?\d+)?", text):
        raise ValueError("Table values must be bounded decimal or rational numbers, not formulas/exponents")
    value = Fraction(text)
    if value.numerator.bit_length() > 8192 or value.denominator.bit_length() > 8192:
        raise ValueError("Numeric result exceeds limit")
    return value


def calculate(expression):
    if not isinstance(expression, str) or len(expression) > 300:
        raise ValueError("Expression must be at most 300 characters")
    source = expression.replace("^", "**")
    if re.search(r"\d[eE][+-]?\d", source):
        raise ValueError("Use decimal literals, not scientific notation")
    tree = ast.parse(source, mode="eval")
    if len(list(ast.walk(tree))) > 100:
        raise ValueError("Expression is too complex")

    def evaluate(node):
        if isinstance(node, ast.Constant) and type(node.value) in (int, float):
            value = Fraction(ast.get_source_segment(source, node))
        elif isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = evaluate(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
        elif isinstance(node, ast.BinOp):
            left, right = evaluate(node.left), evaluate(node.right)
            if isinstance(node.op, ast.Add):
                value = left + right
            elif isinstance(node.op, ast.Sub):
                value = left - right
            elif isinstance(node.op, ast.Mult):
                value = left * right
            elif isinstance(node.op, ast.Div):
                value = left / right
            elif isinstance(node.op, ast.Mod):
                value = left % right
            elif isinstance(node.op, ast.Pow) and right.denominator == 1 and abs(right) <= 1024:
                value = left ** int(right)
            else:
                raise ValueError("Allowed operators: + - * / % and integer powers up to 1024")
        else:
            raise ValueError("Only numeric arithmetic is allowed")
        if value.numerator.bit_length() > 8192 or value.denominator.bit_length() > 8192:
            raise ValueError("Result exceeds limit")
        return value
    return str(evaluate(tree.body))


class Workspace:
    def __init__(self):
        self.token = secrets.token_hex(32)
        self.touched = time.monotonic()
        self.files, self.original = {}, {}
        self.busy = False
        self.pdf_checked = set()
        self.pdf_outcomes = {}
        self.history = []
        self.tool_events = []
        self.active_task = None
        self.active_turn = None
        self.closed = False

    def put(self, name, data, original=False):
        name = filename(name)
        if len(self.files) >= 100 and name not in self.files:
            raise ValueError("Workspace has at most 100 files")
        total = sum(map(len, self.files.values())) - len(self.files.get(name, b"")) + len(data)
        if total > LIMIT:
            raise ValueError("Workspace has at most 8 MiB")
        self.files[name] = data
        if original and name not in self.original:
            self.original[name] = data

    def upload(self, name, data):
        if name.lower().endswith(".zip"):
            try:
                members = unzip(data)
            except zipfile.BadZipFile as exc:
                raise ValueError("Invalid ZIP archive") from exc
            clean = {filename(n): d for n, d in members.items()}
            if self.files.keys() & clean.keys():
                raise ValueError("Upload would overwrite existing files; use distinct project paths")
            if (len(self.files.keys() | clean.keys()) > 100
                    or sum(len(d) for n, d in self.files.items() if n not in clean)
                    + sum(map(len, clean.values())) > LIMIT):
                raise ValueError("Workspace limits exceeded")
            for n, d in clean.items():
                self.put(n, d, original=True)
            return list(clean)
        if name in self.files:
            raise ValueError("Upload would overwrite an existing file; rename before uploading")
        self.put(name, data, original=True)
        return [name]

    def get(self, path):
        name = filename(path)
        if name not in self.files:
            raise ValueError("File not found")
        return self.files[name]

    def read_pdf(self, path, start_page=1, page_count=3, offset=0):
        from pdf_documents import read_pdf
        data = self.get(path)
        return read_pdf(data, path, start_page, page_count, offset)

    def tools(self):
        from agent_framework import tool

        @tool
        def calculator(expression: str) -> str:
            """Exact rational arithmetic (+ - * / % ^); no floating point rounding."""
            return json.dumps({"expression": expression, "result": calculate(expression)})

        @tool
        def get_current_time(timezone: str = "Asia/Seoul") -> str:
            """Current time in an IANA timezone, including UTC offset."""
            return json.dumps({"timezone": timezone, "time": datetime.now(ZoneInfo(timezone)).isoformat()})

        @tool
        def list_files() -> str:
            """List only this chat's files, with sizes."""
            return json.dumps({n: len(d) for n, d in self.files.items()}, ensure_ascii=False)

        @tool
        def read_file(path: str, start: int = 1, count: int = 80) -> str:
            """Read TXT/MD/code, DOCX (paragraph refs), CSV/XLSX. PDF is supported: use read_pdf for page bounds.
            Default read_file on PDF returns the first 3 pages with typed outcomes, never guesses scan from size."""
            if start < 1 or not 1 <= count <= 200:
                raise ValueError("start >= 1 and count 1..200 required")
            if path.lower().endswith(".pdf"):
                if start != 1 or count != 80:
                    raise ValueError("For PDF pagination use read_pdf(start_page, page_count, offset)")
                return json.dumps(self.read_pdf(path), ensure_ascii=False)
            text = extract(path, self.get(path))
            return "\n".join(text.splitlines()[start - 1:start - 1 + count])[:30_000]

        @tool
        def read_pdf(path: str, start_page: int = 1, page_count: int = 3, offset: int = 0) -> str:
            """Read an uploaded PDF using its actual native text layer, including Korean and complex layouts.
            Only this chat's exact file path; 1-based pages, 1..5 pages/call, 10000 chars/page chunk.
            Returns page citations, total_pages, next_page/next_offset and typed outcomes.
            Image-only/blank, corrupt, encrypted and resource-limit results are distinct.
            Never infer unsupported/scanned from filename/size. OCR is not configured; no visual chart analysis."""
            return json.dumps(self.read_pdf(path, start_page, page_count, offset), ensure_ascii=False)

        @tool
        def ocr_pdf(path: str, start_page: int = 1, page_count: int = 3) -> str:
            """OCR of image-only PDF pages is unavailable: no approved Azure Document Intelligence resource.
            No document is sent externally, no billable OCR call is made and no output is invented."""
            from pdf_documents import PDFProblem
            self.get(path)
            if not path.lower().endswith(".pdf"):
                raise PDFProblem("pdf_unsupported", "OCR path accepts PDF files only.", path)
            if start_page < 1 or not 1 <= page_count <= 5:
                raise PDFProblem("pdf_range", "Use start_page >= 1 and page_count 1..5.", path)
            raise PDFProblem("ocr_not_configured", "Azure AI Document Intelligence OCR is not configured. "
                             "An approved existing resource endpoint, authorized managed identity, "
                             "and document-transfer/page-charge approval are required. No OCR was performed.", path,
                             method="ocr_not_performed")

        @tool
        def search_files(query: str) -> str:
            """Literal case-insensitive local text search with file and line references; never web search."""
            if not 1 <= len(query) <= 200:
                raise ValueError("Search query must be 1..200 characters")
            found = []
            for name, data in self.files.items():
                for index, line in enumerate(extract(name, data).splitlines(), 1):
                    if query.casefold() in line.casefold():
                        found.append({"file": name, "line": index, "text": line[:500]})
                        if len(found) == 30:
                            return json.dumps(found, ensure_ascii=False)
            return json.dumps(found, ensure_ascii=False)

        @tool
        def write_file(path: str, content: str) -> str:
            """Create/replace a UTF-8 text project file in this chat. Does not execute code."""
            if PurePosixPath(path).suffix.lower() in {".pdf", ".docx", ".xlsx"}:
                raise ValueError("write_file writes text only")
            if len(content) > TEXT_LIMIT:
                raise ValueError("Text exceeds limit")
            self.put(path, content.encode("utf-8"))
            return json.dumps({"file": path, "bytes": len(content.encode("utf-8")),
                               "artifact": {"kind": "download", "path": path}})

        @tool
        def diff_file(path: str) -> str:
            """Unified diff against the original upload (or empty for a new file)."""
            before = self.original.get(filename(path), b"").decode("utf-8")
            after = self.get(path).decode("utf-8")
            return "".join(difflib.unified_diff(before.splitlines(True), after.splitlines(True),
                                              fromfile="original/" + path, tofile="modified/" + path))[:30_000]

        @tool
        def analyze_table(path: str, operation: str = "profile", column: str = "") -> str:
            """CSV/XLSX first sheet: profile, sum, mean, min, max; exact rational numeric results. No formulas."""
            rows = table_rows(path, self.get(path))
            if not rows:
                raise ValueError("Empty table")
            if operation == "profile":
                return json.dumps({"columns": rows[0], "rows": len(rows) - 1, "sample": rows[1:6]},
                                  ensure_ascii=False)
            if operation not in {"sum", "mean", "min", "max"} or column not in rows[0]:
                raise ValueError("Choose sum/mean/min/max and an exact column header")
            index = rows[0].index(column)
            values = [number(row[index].strip()) for row in rows[1:] if len(row) > index and row[index].strip()]
            if not values:
                raise ValueError("No numeric values")
            total = Fraction(0)
            for item in values:
                total += item
                if total.numerator.bit_length() > 8192 or total.denominator.bit_length() > 8192:
                    raise ValueError("Aggregate numeric result exceeds limit")
            value = {"sum": total, "mean": total / len(values), "min": min(values), "max": max(values)}[operation]
            return json.dumps({"column": column, "operation": operation, "count": len(values), "result": str(value)})

        @tool
        def chart_table(path: str, label_column: str, value_column: str) -> str:
            """Produce an on-demand static SVG bar chart artifact from at most 30 CSV/XLSX rows."""
            import html
            rows = table_rows(path, self.get(path))
            if not rows or label_column not in rows[0] or value_column not in rows[0]:
                raise ValueError("Exact column headers required")
            li, vi = rows[0].index(label_column), rows[0].index(value_column)
            pairs = [(r[li], number(r[vi].strip())) for r in rows[1:31] if len(r) > max(li, vi)]
            if not pairs or any(v < 0 for _, v in pairs):
                raise ValueError("Chart requires non-negative numeric values")
            maximum = max(v for _, v in pairs) or 1
            svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="640" height="{len(pairs)*28+30}">']
            for i, (label, value) in enumerate(pairs):
                y = i * 28 + 20
                svg.append(f'<text x="5" y="{y}" fill="black">{html.escape(label[:20])}</text>'
                           f'<rect x="180" y="{y-15}" width="{float(value/maximum)*350:.2f}" height="20" fill="#568"/>'
                           f'<text x="535" y="{y}" fill="black">{html.escape(str(value))}</text>')
            svg.append("</svg>")
            # SVG is returned as data, never accepted as active uploaded content.
            return json.dumps({"artifact": {"kind": "chart", "svg": "".join(svg)}})

        @tool
        def preview_html(path: str) -> str:
            """Offer an on-demand HTML preview. Browser sandbox disables scripts, forms, navigation and network."""
            if not path.lower().endswith(".html"):
                raise ValueError("Preview requires an HTML file")
            self.get(path).decode("utf-8")
            return json.dumps({"artifact": {"kind": "preview", "path": path}})

        @tool
        def run_tests(command: str) -> str:
            """External isolated test sandbox: unavailable. Never executes uploaded/generated code on frontend or GPU."""
            raise ValueError("Sandbox is not configured: approved isolated execution service with resource/time/network/credential limits required. No tests were run.")

        return [calculator, get_current_time, list_files, read_file, read_pdf, ocr_pdf, search_files, write_file, diff_file,
                analyze_table, chart_table, preview_html, run_tests]


class Workspaces:
    ttl = 1800

    def __init__(self):
        self.items = {}

    def prune(self):
        now = time.monotonic()
        for token, space in list(self.items.items()):
            if not space.busy and now - space.touched > self.ttl:
                del self.items[token]

    def create(self):
        self.prune()
        if len(self.items) >= 16:
            raise ValueError("Workspace capacity reached; try again after existing chats expire")
        space = Workspace()
        self.items[space.token] = space
        return space

    def get(self, token):
        self.prune()
        space = self.items.get(token)
        if space is None:
            raise ValueError("Chat workspace expired; start a new chat")
        space.touched = time.monotonic()
        return space
